"""成绩导出（`GET /api/v1/grades/export` 的**服务层**；路由注册由 lead 统一完成）。

四种关注点各自独立，且都只作用于**同一批行**：

1. **口径一致**：导出不另写 SQL，而是复用 `GradeQueryService.iter_grades` 的同一游标遍历，
   因此"导出口径 = 查询口径"是结构性的（同一 WHERE、同一排序、同一 `GradeRow` 映射），而不是
   两处代码碰巧一致。多页导出与分页查询逐页等价。
2. **编码**：CSV 为 **UTF-8 + BOM**（`\\ufeff`）+ CRLF（RFC 4180）。BOM 是"中文不乱码"的落点：
   没有它，Excel 在中文 Windows 上会按 GBK 解码，表头与姓名变乱码。XLSX 自带 UTF-8 XML，无需 BOM。
3. **防公式注入**：以 `=`、`+`、`-`、`@`、制表符（`\\t`）、回车（`\\r`）开头的单元格值统一前缀
   单引号 `'`（OWASP 建议的转义）。这不是"洁癖"：`openpyxl` 把以 `=` 开头的字符串**写成公式单元格**
   （`data_type == "f"`），未经转义的导出就是把学生的姓名变成一个可执行的公式。
4. **确定性**：同一批行渲染两次必须**逐字节相同**（XLSX 亦然）。这是一条内容性质而不只是性能性质：
   导出文件被下游按内容摘要比对（教师两次下载要看"是不是同一份"、幂等重放要能判定"同一产物"）时，
   一个每次都变的产物等于永远"变了"。`openpyxl` 默认会把**墙钟时间**写进产物两处
   （`docProps/core.xml` 的 `dcterms:created`/`dcterms:modified`、以及 ZIP 条目时间戳），
   因此本模块显式覆盖这两处（见 `_write_xlsx` / `_repack_deterministically`）。

导出内容只有契约 `GradeRow` 的字段：**不含**源码、参考答案、隐藏用例、`report` JSON（报告与逐用例
证据是另一个受权端点的事），也不含内部路径。列顺序即 `GRADE_EXPORT_COLUMNS` 的顺序。
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter
from sqlalchemy import Engine

from autograder_b2.api.formatting import rfc3339_millis
from autograder_b2.auth.tokens import Principal
from autograder_b2.grades.query import (
    DEFAULT_BATCH_SIZE,
    GradeQueryService,
    GradeRow,
)

ExportFormat = Literal["xlsx", "csv"]

#: 契约 `exportGrades` 的 `format` 取值。
EXPORT_FORMATS: tuple[str, ...] = ("xlsx", "csv")

#: 触发公式注入的前导字符（契约明文列举的六个）。
FORMULA_TRIGGERS: tuple[str, ...] = ("=", "+", "-", "@", "\t", "\r")
#: 转义前缀：单引号让 Excel/Calc 把内容当**文本**而不是公式。
ESCAPE_PREFIX = "'"

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv"
CSV_BOM = "\ufeff"

#: 导出列：`(GradeRow 字段, 表头)`。表头用中文（面向教师），字段名照抄契约。
#: 状态列写**状态机原值**（`completed` 而不是"已完成"）：导出必须能与 API 的行逐字段对齐，
#: 加一张翻译表就等于多一条会漂移的口径。
GRADE_EXPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ("submission_id", "提交ID"),
    ("student_id", "学生ID"),
    ("student_number", "学号"),
    ("display_name", "姓名"),
    ("course_id", "课程ID"),
    ("section_id", "课序ID"),
    ("question_version_id", "题目版本ID"),
    ("state", "状态"),
    ("score", "得分"),
    ("max_score", "满分"),
    ("is_pass", "是否通过"),
    ("submitted_at", "提交时间(UTC)"),
    ("evaluation_run_id", "评测运行ID"),
)

#: 时间列的表头（XLSX 写入的日期时间不带时区，Excel 的序列号本来就没有时区）。
_TIMESTAMP_FIELD = "submitted_at"
_TIMESTAMP_NUMBER_FORMAT = "yyyy-mm-dd hh:mm:ss"
_MIN_WIDTH = 10
_MAX_WIDTH = 40

#: 确定性产物的两个时间基准。**都不读墙钟**：
#: - ZIP 条目时间戳固定在 DOS 时间戳的下界（1980-01-01）——ZIP 格式本身也写不下更早的时刻；
#: - 文档属性（`dcterms:created`/`dcterms:modified`）从**行数据**派生（这批成绩里最新的提交时刻），
#:   空批次回落到同一个下界。
_ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)
#: 朴素 UTC（不带时区）：`openpyxl` 的 W3CDTF 序列化器直接输出字段值再补 `Z`，不做时区换算。
_EMPTY_BATCH_MOMENT = datetime(1980, 1, 1, tzinfo=UTC).replace(tzinfo=None)
#: `create_system = 0`（MS-DOS）：不把宿主平台（Windows=0 / Unix=3）写进产物，跨平台产物同字节。
_ZIP_CREATE_SYSTEM = 0
#: 与 `zipfile.ZipFile.writestr` 对普通文件设置的权限位保持一致（0o600）。
_ZIP_EXTERNAL_ATTR = 0o600 << 16


class UnsupportedExportFormat(ValueError):
    """`format` 不在 `xlsx`/`csv` 内 → 对外 422（路由层也可用枚举声明在解析阶段拦下）。"""


@dataclass(frozen=True, slots=True)
class ExportFile:
    """一个可直接作为 HTTP 响应体返回的导出文件。"""

    content: bytes
    media_type: str
    filename: str
    format: str
    row_count: int

    @property
    def content_disposition(self) -> str:
        """`attachment; filename="..."`；文件名为纯 ASCII，因此无需 RFC 5987 编码。"""
        return f'attachment; filename="{self.filename}"'


def escape_spreadsheet_cell(value: str) -> str:
    """公式注入转义：以触发字符开头的值前缀 `'`；其余原样返回。

    幂等：结果以 `'` 开头，而 `'` 不在触发集合内，因此重复调用不会叠加前缀。
    `-` 也在集合内（契约明文列举）：`-2+3+cmd|' /C calc'!A0` 这类以减号开头的 DDE 载荷同样危险。
    数字单元格（`score`/`max_score`）写原生数值类型、不经本函数：Excel 不会把数值当公式，
    而把一个数变成文本会破坏求和与排序。
    """
    if value.startswith(FORMULA_TRIGGERS):
        return ESCAPE_PREFIX + value
    return value


def render_csv(rows: Sequence[GradeRow]) -> bytes:
    """渲染 CSV：UTF-8 BOM + CRLF + 公式注入转义。

    空值（未完成运行的分数）写成**空字段**而不是 `0`：表里没有值就是没有值，`0` 是一个结论。
    """
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([header for _, header in GRADE_EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([_csv_text(row, field) for field, _ in GRADE_EXPORT_COLUMNS])
    return (CSV_BOM + buffer.getvalue()).encode("utf-8")


def render_xlsx(rows: Sequence[GradeRow]) -> bytes:
    """渲染 XLSX（单工作表，首行加粗冻结）。

    数值列写原生数值、布尔列写原生布尔、时间列写 UTC 的**朴素** `datetime`（Excel 日期序列号没有
    时区概念，表头标注 `(UTC)` 以免被当成本地时间）。空值跳过写入 → 单元格为**空**而不是 `0`。

    **确定性**：同一批行两次调用的返回值逐字节相同。`openpyxl` 会把墙钟时间写进 `docProps/core.xml`
    与每个 ZIP 条目，因此这里显式盖上从行数据派生的文档时间（`_document_moment`），并把包重打成固定
    条目时间戳的 ZIP（`_repack_deterministically`）。OOXML 结构、数字格式、列宽、冻结窗格都不受影响。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩"
    width_units = [_display_width(header) for _, header in GRADE_EXPORT_COLUMNS]
    for column, (_, header) in enumerate(GRADE_EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
    for index, row in enumerate(rows, start=2):
        for column, (field, _) in enumerate(GRADE_EXPORT_COLUMNS, start=1):
            value = _xlsx_value(row, field)
            if value is None:
                continue
            cell = sheet.cell(row=index, column=column, value=value)
            if field == _TIMESTAMP_FIELD:
                cell.number_format = _TIMESTAMP_NUMBER_FORMAT
            width_units[column - 1] = max(width_units[column - 1], _display_width(_cell_text(value)))
    sheet.freeze_panes = "A2"
    for column, units in enumerate(width_units, start=1):
        width = min(_MAX_WIDTH, max(_MIN_WIDTH, units + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width
    # 不需要 `workbook.close()`：写出走的是本函数自己持有并关闭的 ZIP 归档（不是工作簿的文件句柄）。
    return _write_xlsx(workbook, _document_moment(rows))


def _document_moment(rows: Sequence[GradeRow]) -> datetime:
    """本次产物的文档时间：从**行数据**派生（这批行里最新的 `submitted_at`，截到秒，UTC 朴素）。

    为什么不是常量：产物的元数据仍然应该说明点什么（"这批成绩新到什么时刻"），而墙钟只说明
    "服务器什么时候被调用"，与内容无关。为什么不是墙钟：见 `_write_xlsx`。
    截到秒是因为 `openpyxl` 的 W3CDTF 序列化器按 `timespec="seconds"` 输出（微秒写不进去）。
    **必须是 UTC 朴素值**：该序列化器直接输出字段值再补 `Z`，不做时区换算。
    空批次（比如某门课还没有提交）用固定基准 `_EMPTY_BATCH_MOMENT`。
    """
    moment = _EMPTY_BATCH_MOMENT
    for row in rows:
        candidate = _xlsx_moment(row.submitted_at).replace(microsecond=0)
        moment = max(moment, candidate)
    return moment


def _write_xlsx(workbook: Workbook, moment: datetime) -> bytes:
    """写出 XLSX：文档时间由调用方给定，ZIP 条目时间戳固定。

    这里**不**调用 `openpyxl.writer.excel.save_workbook`：它在写盘前把
    `workbook.properties.modified` 覆盖成 `datetime.now(utc)`（已实测），于是"同一批行 → 不同字节"。
    改用 `save_workbook` 内部用的同一个写出器 `ExcelWriter`，把两个时间都钉死；ZIP 条目时间戳
    （`ZipFile.writestr` 默认取 `time.localtime()`）由 `_repack_deterministically` 统一重打。
    """
    workbook.properties.created = moment
    workbook.properties.modified = moment
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(workbook, archive).save()
    return _repack_deterministically(buffer.getvalue())


def _repack_deterministically(payload: bytes) -> bytes:
    """把 `openpyxl` 交出的包重打成**确定性** ZIP（条目时间戳固定、条目顺序与压缩方式不变）。

    为什么必须重打：`zipfile.ZipFile.writestr` 默认用 `time.localtime()` 当每个条目的时间戳，
    于是"同一输入 → 不同字节"。XLSX 是 ZIP，字节一变，按内容摘要做的比对（幂等重放、去重、
    "两次下载是不是同一份"）全部失真，测试也会变成天然 flaky。重打只动时间戳与 `create_system`，
    条目内容、顺序、压缩方式、以及 OOXML 结构都不变，Excel 照常打开。
    """
    source = zipfile.ZipFile(io.BytesIO(payload))
    target = io.BytesIO()
    try:
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            for info in source.infolist():
                entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_DATE_TIME)
                entry.compress_type = zipfile.ZIP_DEFLATED
                entry.create_system = _ZIP_CREATE_SYSTEM
                entry.external_attr = _ZIP_EXTERNAL_ATTR
                entry.internal_attr = info.internal_attr
                archive.writestr(entry, source.read(info.filename))
    finally:
        source.close()
    return target.getvalue()


class GradesExportService:
    """成绩导出（`exportGrades` 的服务层）；取数完全委托 `GradeQueryService`。"""

    def __init__(self, engine: Engine, *, batch_size: int = DEFAULT_BATCH_SIZE) -> None:
        self._query = GradeQueryService(engine, batch_size=batch_size)

    def export(
        self,
        *,
        principal: Principal,
        course_id: str,
        section_id: str | None = None,
        question_version_id: str | None = None,
        student_id: str | None = None,
        fmt: str = "xlsx",
    ) -> ExportFile:
        """导出成绩；过滤语义与 `list_grades` **同一个** `_resolve`（含授权与课序收敛）。

        参数 `fmt` 而非 `format`：路由层的查询参数叫 `format`（契约），这里避开同名内建函数。
        异常与查询一致：`AuthorizationError` → 403，`InvalidFilterError`/`InvalidCursorError`
        → 422，`UnsupportedExportFormat` → 422。
        """
        if fmt not in EXPORT_FORMATS:
            raise UnsupportedExportFormat(f"format must be one of {', '.join(EXPORT_FORMATS)}")
        rows = list(
            self._query.iter_grades(
                principal=principal,
                course_id=course_id,
                section_id=section_id,
                question_version_id=question_version_id,
                student_id=student_id,
            )
        )
        content = render_xlsx(rows) if fmt == "xlsx" else render_csv(rows)
        extension = "xlsx" if fmt == "xlsx" else "csv"
        return ExportFile(
            content=content,
            media_type=XLSX_MEDIA_TYPE if fmt == "xlsx" else CSV_MEDIA_TYPE,
            filename=f"grades-{course_id}.{extension}",
            format=fmt,
            row_count=len(rows),
        )


def _csv_text(row: GradeRow, field: str) -> str:
    """CSV 单元格文本（统一口径，与格式无关）：空值→空串，布尔→`true`/`false`。"""
    value = getattr(row, field)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return rfc3339_millis(value)
    if isinstance(value, (int, float)):
        return _decimal_text(float(value))
    return escape_spreadsheet_cell(str(value))


def _xlsx_value(row: GradeRow, field: str) -> object | None:
    """XLSX 单元格值：原生类型优先；文本经公式注入转义。"""
    value = getattr(row, field)
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, datetime):
        return _xlsx_moment(value)
    return escape_spreadsheet_cell(str(value))


def _xlsx_moment(value: datetime) -> datetime:
    """XLSX 时间单元格：UTC、**朴素**（Excel 序列号无时区），且**截断到毫秒**。

    截断而不是交给 openpyxl：`rfc3339_millis` 对外是毫秒**截断**，而 openpyxl 读写 Excel 序列号时
    会四舍五入到毫秒——对**带微秒的来源**，`.123999` 在 CSV/JSON 里是 `.123`、在 XLSX 里会变成
    `.124`，同一时刻在两种格式里差 1 毫秒，正是"导出口径一致"要防的那类不一致。

    今天从库里读到的 `submitted_at` 已经是毫秒对齐的（`timestamptz(3)`，迁移 0012），因此这一步
    对它是**幂等**的；留着它是因为它守的是导出层的不变量而不是当前列宽：来源一旦带上更高精度
    （合成行，或将来放宽列精度），两种格式就又会分叉。直接验证这条不变量的用例是
    `tests/test_grades_query_export.py::test_xlsx_truncates_like_the_api_when_the_source_carries_microseconds`
    ——它用**合成行**注入微秒，因为库列已经不可能存微秒。
    """
    moment = value.astimezone(UTC) if value.tzinfo else value.replace(tzinfo=UTC)
    return moment.replace(microsecond=(moment.microsecond // 1000) * 1000, tzinfo=None)


def _decimal_text(value: float) -> str:
    """`float` → 十进制文本，去掉多余的尾随零（`86.0` → `86`，`86.50` → `86.5`）。

    经 `Decimal(str(...))` 而不是 `%f`/`round`：后者会引入二进制浮点的近似尾巴。
    """
    text = format(Decimal(str(value)), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _cell_text(value: object) -> str:
    """列宽估算用的显示文本。"""
    if isinstance(value, datetime):
        return rfc3339_millis(value)
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return _decimal_text(value)
    return str(value)


def _display_width(text: str) -> int:
    """粗估显示宽度：CJK/全角字符按 2 列计（Excel 列宽以"字符数"为单位）。"""
    return sum(2 if ord(char) > 0x2E80 else 1 for char in text)
